"""表格读写：支持 xlsx（优先）与 csv。

- 读：pandas，自动识别 sheet
- 写：openpyxl 原地写列，保留其他 sheet 与格式，避免整表被覆盖
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")

# 纯文本表格：后缀 -> 分隔符。必须与 core.source.TABLE_EXTS 保持一致，
# 否则一个「支持 .tsv」的来源会在这里被当成 Excel 去解析而崩溃。
TEXT_DELIM = {".csv": ",", ".tsv": "\t"}


def _is_text_table(p: Path) -> bool:
    return p.suffix.lower() in TEXT_DELIM


def read_table(path: Path | str, sheet: str | None = None) -> tuple[list[dict], list[str]]:
    """返回 (行列表, 列名列表)。值统一转成字符串，空值转空串。"""
    p = Path(path)
    if _is_text_table(p):
        df = pd.read_csv(p, dtype=str, sep=TEXT_DELIM[p.suffix.lower()]).fillna("")
    else:
        if sheet:
            df = pd.read_excel(p, sheet_name=sheet, dtype=str).fillna("")
        else:
            df = pd.read_excel(p, dtype=str).fillna("")
    df = df.astype(str)
    columns = [str(c) for c in df.columns]
    rows = [{c: str(r[i]) for i, c in enumerate(columns)} for r in df.itertuples(index=False)]
    return rows, columns


def list_sheets(path: Path | str) -> list[str]:
    p = Path(path)
    if _is_text_table(p):
        # 文本表格没有 sheet 概念，返回后缀名作为标识即可
        return [p.suffix.lower().lstrip(".")]
    return load_workbook(p, read_only=True).sheetnames


def resolve_column(columns: list[str], spec: str) -> str:
    """按列名或 1 基序号定位列；支持「包含」匹配。"""
    spec = str(spec).strip()
    if spec.isdigit():
        idx = int(spec) - 1
        if 0 <= idx < len(columns):
            return columns[idx]
        raise KeyError(f"列序号越界：{spec}，共 {len(columns)} 列")
    if spec in columns:
        return spec
    hits = [c for c in columns if spec in c]
    if len(hits) == 1:
        return hits[0]
    if len(hits) > 1:
        raise KeyError(f"列名 {spec} 匹配到多个：{hits}，请写完整列名")
    raise KeyError(f"找不到列 {spec}，可用列：{columns}")


def write_column(path: Path | str, column: str, values: list[Any],
                 sheet: str | None = None, key_column: str | None = None) -> Path:
    """把一列写回表格。

    - key_column 为空：按行序写入
    - key_column 指定：按该列的值匹配行写入（乱序安全）
    """
    p = Path(path)
    if _is_text_table(p):
        sep = TEXT_DELIM[p.suffix.lower()]
        df = pd.read_csv(p, dtype=str, sep=sep).fillna("")
        if column not in df.columns:
            df[column] = ""
        if key_column and key_column in df.columns:
            # values 形如 [(key, value)]
            mapping = {str(k): v for k, v in values}
            df[column] = df[key_column].astype(str).map(mapping).fillna(df[column])
        else:
            df[column] = [v for _, v in values] if values and isinstance(values[0], tuple) else list(values)
        df.to_csv(p, index=False, encoding="utf-8-sig", sep=sep)
        return p

    wb = load_workbook(p)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = [c.value for c in ws[1]]
    headers = [str(h) if h is not None else "" for h in headers]

    if column in headers:
        col_idx = headers.index(column) + 1
    else:
        col_idx = len(headers) + 1
        ws.cell(row=1, column=col_idx, value=column)
        ws.cell(row=1, column=col_idx).font = Font(bold=True)
        ws.cell(row=1, column=col_idx).fill = HEADER_FILL
        headers.append(column)

    if key_column and key_column in headers:
        key_idx = headers.index(key_column) + 1
        row_of: dict[str, int] = {}
        for r in range(2, ws.max_row + 1):
            v = ws.cell(row=r, column=key_idx).value
            if v not in (None, ""):
                row_of.setdefault(str(v), r)
        for key, val in values:
            r = row_of.get(str(key))
            if r:
                c = ws.cell(row=r, column=col_idx, value=val)
                c.alignment = Alignment(wrap_text=True, vertical="top")
    else:
        for i, item in enumerate(values):
            # 调用方常传 (key, value) 元组，写单元格前先取到真正的文本
            val = item[1] if isinstance(item, (tuple, list)) and len(item) == 2 else item
            c = ws.cell(row=2 + i, column=col_idx, value=val)
            c.alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 60
    wb.save(p)
    return p


def export_rows(path: Path | str, rows: list[dict], columns: list[str] | None = None,
                sheet: str = "Sheet1") -> Path:
    """新建表格文件并写入全部行。"""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    cols = columns or (list(rows[0].keys()) if rows else [])

    if _is_text_table(p):
        pd.DataFrame(rows, columns=cols).to_csv(
            p, index=False, encoding="utf-8-sig", sep=TEXT_DELIM[p.suffix.lower()])
        return p

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for i, row in enumerate(rows, 2):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=i, column=j, value=row.get(c, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = \
            max(12, min(60, len(str(c)) * 2 + 10))
    wb.save(p)
    return p
